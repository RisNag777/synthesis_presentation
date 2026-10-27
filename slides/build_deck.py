"""Redesign Beyond the Label as a single visual system."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

INK = RGBColor(0x1C, 0x19, 0x17)
MUTED = RGBColor(0x5C, 0x57, 0x4F)
FAINT = RGBColor(0x8A, 0x84, 0x7C)
PAPER = RGBColor(0xF6, 0xF4, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x1F, 0x4E, 0x5F)
SLATE_SOFT = RGBColor(0xE6, 0xEE, 0xF0)
CLAY = RGBColor(0x8C, 0x4A, 0x32)
CLAY_SOFT = RGBColor(0xF3, 0xEB, 0xE6)
TRACK = RGBColor(0xE4, 0xE0, 0xDA)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xDD, 0xD8, 0xD2)

W = Inches(13.333)
H = Inches(7.5)
FONT = "Calibri"


def strip_shadow(shape):
    sp_pr = shape._element.spPr
    effect = sp_pr.find(qn("a:effectLst"))
    if effect is not None:
        sp_pr.remove(effect)


def rect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    strip_shadow(shape)
    return shape


def rrect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    strip_shadow(shape)
    return shape


def text(slide, l, t, w, h, content, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(l, t, w, h)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    lines = content.split("\n") if isinstance(content, str) else content
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = FONT
    return shape


def bg(slide):
    rect(slide, 0, 0, W, H, PAPER)


def eyebrow(slide, label):
    text(slide, Inches(0.62), Inches(0.36), Inches(10), Inches(0.28), label.upper(), 12, SLATE, bold=True)


def title(slide, content, top=0.68, size=32, height=1.15):
    return text(slide, Inches(0.62), Inches(top), Inches(12.1), Inches(height), content, size, INK, bold=True)


def new(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    return slide


def card_title_body(slide, l, t, w, h, fill, kicker, heading, body, kicker_color=SLATE):
    rrect(slide, l, t, w, h, fill)
    text(slide, l + Inches(0.28), t + Inches(0.26), w - Inches(0.5), Inches(0.28), kicker, 13, kicker_color, bold=True)
    text(slide, l + Inches(0.28), t + Inches(0.58), w - Inches(0.5), Inches(0.85), heading, 20, INK, bold=True)
    text(slide, l + Inches(0.28), t + Inches(1.5), w - Inches(0.5), h - Inches(1.75), body, 15, MUTED)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Beyond the Label"
    prs.core_properties.subject = "Restorative alternatives to discipline in middle school"

    # 1 Title
    s = new(prs)
    rect(s, 0, 0, Inches(0.18), H, SLATE)
    text(s, Inches(0.85), Inches(1.85), Inches(11), Inches(0.3), "WASHINGTON MIDDLE SCHOOLS", 14, SLATE, bold=True)
    text(s, Inches(0.85), Inches(2.25), Inches(11.5), Inches(1.5), "Beyond the label", 60, INK, bold=True)
    text(
        s,
        Inches(0.85),
        Inches(3.9),
        Inches(10.5),
        Inches(1.1),
        "Restorative alternatives to discipline in middle school.",
        24,
        MUTED,
    )
    rect(s, Inches(0.85), Inches(5.25), Inches(2.1), Inches(0.06), SLATE)
    text(
        s,
        Inches(0.85),
        Inches(5.5),
        Inches(10),
        Inches(0.7),
        "From pathologizing adolescent distress\nto classrooms built for repair and belonging.",
        16,
        MUTED,
    )

    # 2 Context
    s = new(prs)
    eyebrow(s, "Context  ·  2023–24")
    title(s, "Students with disabilities are 15% of enrollment\nand 89% of restraint and isolation incidents.", size=30, height=1.25)

    def bar(top, label, detail, share, inside):
        text(s, Inches(0.62), top, Inches(7.6), Inches(0.3), label, 15, INK, bold=True)
        text(s, Inches(8.2), top, Inches(4.5), Inches(0.3), detail, 15, MUTED, align=PP_ALIGN.RIGHT)
        bt = top + Inches(0.38)
        bw = Inches(12.05)
        rect(s, Inches(0.62), bt, bw, Inches(0.38), TRACK)
        fill_w = int(bw * share)
        rect(s, Inches(0.62), bt, fill_w, Inches(0.38), SLATE)
        pct = f"{round(share * 100)}%"
        if inside:
            text(s, Inches(0.78), bt + Inches(0.03), Inches(1.2), Inches(0.32), pct, 14, WHITE, bold=True)
        else:
            text(s, Inches(0.62) + fill_w + Inches(0.12), bt + Inches(0.03), Inches(1.1), Inches(0.32), pct, 14, SLATE, bold=True)

    bar(Inches(2.25), "Share of middle school enrollment", "32,147 of 218,387 students", 32147 / 218387, False)
    bar(Inches(3.45), "Share of restraint and isolation incidents", "1,176 of 1,325 incidents", 1176 / 1325, True)
    text(
        s,
        Inches(0.62),
        Inches(5.15),
        Inches(12.05),
        Inches(1.05),
        "Students without disabilities account for 81 incidents. Another 68 sit in schools where the disability count was suppressed. In 2019–20, when counts of one and two were still published, 97% of middle school restraint incidents and 97% of isolation incidents involved students with disabilities.",
        15,
        MUTED,
    )
    text(
        s,
        Inches(0.62),
        Inches(6.55),
        Inches(12.05),
        Inches(0.6),
        "OSPI, schools with “middle” in the name. “Students with disabilities” is the IDEA category, broader than an emotional or behavioral label. Counts under 3 incidents are suppressed in 2023–24.",
        12,
        FAINT,
    )

    # 3 Why
    s = new(prs)
    eyebrow(s, "The pattern")
    title(s, "Why are disabled students more likely to be suspended?", size=30, height=0.9)
    cards = [
        ("01", "Deficit mislabeling", "Stereotypes and thin training lead educators to read a sensory meltdown as defiance or threat. The behavior is treated as a deficit instead of met with de-escalation, evidence-based support, or a look at the cause."),
        ("02", "Rights going unmet", "IDEA and Section 504 require accommodations, positive behavioral supports, and the safeguards in an IEP or 504 plan. Schools often skip those duties and discipline students in ways that violate their rights."),
        ("03", "Pushout", "Suspension, expulsion, and removal cut disabled students off from learning. Minor incidents are handed to school resource officers, and referrals move students toward the school-to-prison pipeline."),
    ]
    cw = Inches(3.9)
    for i, (k, hdg, body) in enumerate(cards):
        card_title_body(s, Inches(0.62) + i * Inches(4.12), Inches(2.25), cw, Inches(3.85), CARD, k, hdg, body)

    # 4 Brain
    s = new(prs)
    eyebrow(s, "Development")
    title(s, "The middle school brain is built to respond to the room.", size=30, height=0.95)
    brain = [
        ("01", "Plasticity", "Early adolescence is a period of rapid neural remodeling. The nervous system is highly sensitive to stress, perceived threat, and trauma."),
        ("02", "Executive function", "Working memory and cognitive flexibility fluctuate. When academic or sensory demand overloads them, dysregulation is a stress response, not willful defiance."),
        ("03", "Co-regulation", "Because the brain is this responsive, a calm adult in the room can lower fight-or-flight in the moment and build self-regulation over time."),
    ]
    for i, (k, hdg, body) in enumerate(brain):
        card_title_body(s, Inches(0.62) + i * Inches(4.12), Inches(2.3), cw, Inches(3.7), CARD, k, hdg, body)

    # 5 Paradigms
    s = new(prs)
    eyebrow(s, "The shift")
    title(s, "From “what is wrong with this student?”\nto “what does this student need?”", size=30, height=1.2)
    left = Inches(0.62)
    col_w = Inches(5.9)
    top = Inches(2.25)
    col_h = Inches(3.85)
    rrect(s, left, top, col_w, col_h, CLAY_SOFT)
    rrect(s, Inches(6.78), top, col_w, col_h, SLATE_SOFT)
    text(s, left + Inches(0.32), top + Inches(0.28), col_w - Inches(0.6), Inches(0.3), "TRADITIONAL", 13, CLAY, bold=True)
    text(s, Inches(7.1), top + Inches(0.28), col_w - Inches(0.6), Inches(0.3), "RESTORATIVE", 13, SLATE, bold=True)
    text(s, left + Inches(0.32), top + Inches(0.7), col_w - Inches(0.64), Inches(3.5),
         "Dysregulation is read as pathology: emotional disturbance, oppositional defiance.\n\nThe problem is placed in the child: non-compliant, disruptive, unmotivated.\n\nThe response is behavior modification, surveillance, and removal.",
         16, INK)
    text(s, Inches(7.1), top + Inches(0.7), col_w - Inches(0.64), Inches(3.5),
         "Ross Greene: kids do well if they can. Behavior signals lagging skills and an unmet need.\n\nDisCrit: racialized and ableist systems pathologize students. Change the barrier, not the child.\n\nThe response is a changed environment, executive-function support, and repair.",
         16, INK)

    # 6 SAMHSA
    s = new(prs)
    eyebrow(s, "SAMHSA’s trauma-informed approach")
    title(s, "Safety is a condition for learning.", size=32, height=0.7)
    text(
        s,
        Inches(0.62),
        Inches(1.5),
        Inches(12.1),
        Inches(0.7),
        "Chronic stress interrupts focus and self-regulation. Seclusion and restraint re-traumatize. With relational safety and the right support, students can recover.",
        16,
        MUTED,
    )
    rs = [
        ("Realize", "Dysregulation is a survival response, not intentional defiance."),
        ("Recognize", "Learn how trauma shows up by age, gender, and setting."),
        ("Respond", "Put the principle into policy, training, and the design of the room."),
        ("Resist\nre-traumatization", "End seclusion and restraint that replay past harm."),
    ]
    rw = Inches(2.95)
    for i, (hdg, body) in enumerate(rs):
        x = Inches(0.62) + i * Inches(3.15)
        rrect(s, x, Inches(2.7), rw, Inches(3.15), CARD)
        text(s, x + Inches(0.24), Inches(2.95), rw - Inches(0.45), Inches(1.05), hdg, 20, SLATE, bold=True)
        text(s, x + Inches(0.24), Inches(4.1), rw - Inches(0.45), Inches(1.4), body, 16, INK)

    # 7 Practices
    s = new(prs)
    eyebrow(s, "In the room")
    title(s, "Four practices that keep students in the classroom.", size=30, height=0.8)
    practices = [
        ("Predictable routines", "Visual anchors lower the working-memory load and the anxiety of not knowing what comes next."),
        ("Cool-down corners", "A voluntary sensory reset inside the room, so regulation does not require removal."),
        ("The 2×10", "Two minutes of personal connection, for ten days in a row, before a crisis."),
        ("Asset-based notes", "In the IEP, name the unmet need. Do not write the behavior as a character flaw."),
    ]
    pw, ph = Inches(6.0), Inches(2.2)
    for i, (hdg, body) in enumerate(practices):
        x = Inches(0.62) + (i % 2) * Inches(6.22)
        y = Inches(2.15) + (i // 2) * Inches(2.42)
        rrect(s, x, y, pw, ph, CARD)
        text(s, x + Inches(0.3), y + Inches(0.28), pw - Inches(0.55), Inches(0.45), hdg, 20, INK, bold=True)
        text(s, x + Inches(0.3), y + Inches(0.85), pw - Inches(0.55), Inches(1.1), body, 16, MUTED)

    # 8 Policy
    s = new(prs)
    eyebrow(s, "Beyond one classroom")
    title(s, "What the system has to change.", size=32, height=0.7)
    policies = [
        ("01", "End isolation", "Support the Washington effort to eliminate isolation booths."),
        ("02", "Fund the training", "Move discipline dollars into schoolwide restorative practice."),
        ("03", "Call families first", "Learn home regulation strategies before a crisis, not after a write-up."),
        ("04", "Publish the disparity", "Hold districts accountable for discipline rates among disabled students."),
    ]
    for i, (n, hdg, body) in enumerate(policies):
        y = Inches(1.85) + i * Inches(1.28)
        rrect(s, Inches(0.62), y, Inches(12.08), Inches(1.12), CARD)
        text(s, Inches(0.9), y + Inches(0.28), Inches(0.8), Inches(0.55), n, 22, SLATE, bold=True)
        text(s, Inches(1.9), y + Inches(0.18), Inches(10.3), Inches(0.4), hdg, 20, INK, bold=True)
        text(s, Inches(1.9), y + Inches(0.58), Inches(10.3), Inches(0.4), body, 16, MUTED)

    # 9 Families
    s = new(prs)
    eyebrow(s, "Partnership")
    title(s, "Families are experts, not recipients of notices.", size=30, height=0.85)
    fam = [
        ("Before the crisis", "Ask families what already helps their child regulate, and what sets the day off. Treat that knowledge as part of the plan."),
        ("How we write home", "Replace deficit incident reports with collaborative problem-solving. Families should be partners in repair."),
        ("Students at the table", "Middle schoolers help write their sensory-break plans, IEP accommodations, and community agreements."),
    ]
    for i, (hdg, body) in enumerate(fam):
        card_title_body(s, Inches(0.62) + i * Inches(4.12), Inches(2.3), cw, Inches(3.7), CARD, f"0{i+1}", hdg, body)

    # 10 Close
    s = new(prs)
    rect(s, 0, 0, Inches(0.18), H, SLATE)
    text(s, Inches(0.85), Inches(0.55), Inches(11), Inches(0.3), "OUR CHARGE", 14, SLATE, bold=True)
    text(s, Inches(0.85), Inches(1.05), Inches(11.5), Inches(1.2), "Humanize the student\nin front of you.", 44, INK, bold=True)
    charges = [
        ("Dismantle the pipeline", "Stop using a diagnosis to justify seclusion, restraint, and pushout."),
        ("Change the room", "Replace surveillance with executive-function support, de-escalation, and repair."),
        ("Protect belonging", "Guard the dignity and agency of every disabled and neurodivergent student."),
    ]
    for i, (hdg, body) in enumerate(charges):
        y = Inches(3.15) + i * Inches(1.25)
        text(s, Inches(0.85), y, Inches(11), Inches(0.4), hdg, 22, INK, bold=True)
        text(s, Inches(0.85), y + Inches(0.42), Inches(11), Inches(0.45), body, 18, MUTED)

    # 11 References
    s = new(prs)
    eyebrow(s, "Sources")
    title(s, "References", size=32, height=0.6)
    # First appearance: 2 context, 3 title then cards, 4 cards, 5 shift, 6 SAMHSA, 7 2x10, 8 repair.
    refs = [
        [
            ("Office of Superintendent of Public Instruction. (2025). ", False),
            ("2023–24 restraint and isolation data", True),
            (" [Data set].", False),
        ],
        [
            ("The Education Trust. (2020, August 12). ", False),
            ("How school discipline disproportionately targets students with disabilities", True),
            (".", False),
        ],
        [
            ("National Center for Learning Disabilities. (2020). ", False),
            ("Significant disproportionality in special education: Current trends and actions for impact", True),
            (".", False),
        ],
        [
            ("Baglieri, S. (2023). ", False),
            ("Disability studies and the inclusive classroom: Critical practices for embracing diversity in education", True),
            (" (3rd ed.). Routledge.", False),
        ],
        [
            ("Disability Rights Washington & American Civil Liberties Union of Washington. (2023). ", False),
            ("Coming into the light: An examination of restraint and isolation practices in Washington schools", True),
            (".", False),
        ],
        [
            ("Morris, M. W. (2016). ", False),
            ("Pushout: The criminalization of Black girls in schools", True),
            (". The New Press.", False),
        ],
        [
            ("National Child Traumatic Stress Network. (2008). ", False),
            ("Child trauma toolkit for educators", True),
            (".", False),
        ],
        [
            ("Wood, C. (2015). ", False),
            ("Yardsticks: Children in the classroom, ages 4–14", True),
            (" (4th ed.). Center for Responsive Schools.", False),
        ],
        [
            ("Ginwright, S. (2016). ", False),
            ("Hope and healing in urban education: How urban activists and teachers are reclaiming matters of the heart", True),
            (". Routledge.", False),
        ],
        [
            ("Greene, R. W. (2008). Kids do well if they can. ", False),
            ("Phi Delta Kappan, 90", True),
            ("(3), 160–167.", False),
        ],
        [
            ("Greene, R. W. (2014). ", False),
            ("Lost at school: Why our kids with behavioral challenges are falling through the cracks and how we can help them", True),
            (". Scribner.", False),
        ],
        [
            ("Annamma, S. A., Ferri, B. A., & Connor, D. J. (Eds.). (2022). ", False),
            ("DisCrit expanded: Reverberations, ruptures, and inquiries", True),
            (". Teachers College Press.", False),
        ],
        [
            ("Davis, D. R., & Annamma, S. A. (2023). Disabling discipline: DisCrit, restorative justice, and anti-ableist schooling. ", False),
            ("Remedial and Special Education, 44", True),
            ("(1), 18–29.", False),
        ],
        [
            ("Substance Abuse and Mental Health Services Administration. (2014). ", False),
            ("SAMHSA’s concept of trauma and guidance for a trauma-informed approach", True),
            (" (HHS Publication No. SMA 14-4884). U.S. Department of Health and Human Services.", False),
        ],
        [
            ("Mendler, A. N. (2001). ", False),
            ("Connecting with students", True),
            (". Association for Supervision and Curriculum Development.", False),
        ],
        [
            ("Mingus, M. (2019). ", False),
            ("Transformative justice: A brief description", True),
            (". Leaving Evidence.", False),
        ],
    ]

    def ref_block(x, items):
        box = text(s, x, Inches(1.48), Inches(6.05), Inches(5.7), "", 12, INK)
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        for i, parts in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(0)
            p.space_after = Pt(8)
            p.line_spacing = 1.0
            p.level = 0
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.22))))
            pPr.set("indent", str(int(Inches(-0.22))))
            for chunk, italic in parts:
                run = p.add_run()
                run.text = chunk
                run.font.size = Pt(12)
                run.font.color.rgb = INK
                run.font.name = FONT
                run.font.italic = italic

    ref_block(Inches(0.55), refs[:8])
    ref_block(Inches(6.75), refs[8:])
    notes = s.notes_slide.notes_text_frame
    notes.text = (
        "Full entries, including retrieval links. "
        "The Education Trust page dated August 12, 2020 could not be located, so that entry has no URL.\n\n"
        "Office of Superintendent of Public Instruction. (2025). 2023–24 restraint and isolation data [Data set]. "
        "https://ospi.k12.wa.us/student-success/health-safety/school-safety-center/restraint-and-isolation\n\n"
        "National Center for Learning Disabilities. (2020). Significant disproportionality in special education: "
        "Current trends and actions for impact. "
        "https://ncld.org/wp-content/uploads/2025/09/2020-Significant-Disproportionality-in-Special-Education.pdf\n\n"
        "Disability Rights Washington & American Civil Liberties Union of Washington. (2023). Coming into the light: "
        "An examination of restraint and isolation practices in Washington schools. "
        "https://disabilityrightswa.org/wp-content/uploads/2023/02/Restraint-and-Isolation-Report-dist.pdf"
    )
    return prs


if __name__ == "__main__":
    out = r"C:\Users\risha\Documents\Projects\synthesis_presentation\Untitled presentation (2).pptx"
    build().save(out)
    print(out)
